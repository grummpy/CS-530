import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def markdown_to_docx(md_path, docx_path, embedded_images=None):
    doc = Document()
    
    # Configure normal style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_block_lines = []
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        # Determine number of columns
        col_count = max(len(r) for r in table_rows)
        table = doc.add_table(rows=len(table_rows), cols=col_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        
        for r_idx, row_data in enumerate(table_rows):
            row = table.rows[r_idx]
            is_header = (r_idx == 0)
            for c_idx in range(col_count):
                cell = row.cells[c_idx]
                text = row_data[c_idx] if c_idx < len(row_data) else ""
                cell.text = text.strip()
                
                # Format text
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    for run in p.runs:
                        run.font.size = Pt(9.5)
                        if is_header:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                
                if is_header:
                    set_cell_background(cell, "1F497D") # Navy blue
                elif r_idx % 2 == 1:
                    set_cell_background(cell, "F2F2F2") # Light gray alternate
        
        doc.add_paragraph() # Spacing
        table_rows = []
        in_table = False

    def flush_code_block():
        nonlocal in_code_block, code_block_lines
        if not code_block_lines:
            in_code_block = False
            return
        
        code_text = "".join(code_block_lines)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.25)
        
        run = p.add_run(code_text.rstrip())
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
        
        code_block_lines = []
        in_code_block = False

    # Process lines
    for line in lines:
        stripped = line.strip()

        # Handle code fences
        if stripped.startswith("```"):
            if in_code_block:
                flush_code_block()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_block_lines = []
            continue

        if in_code_block:
            code_block_lines.append(line)
            continue

        # Handle Markdown Tables
        if "|" in line and (line.strip().startswith("|") or line.strip().endswith("|")):
            # Check if separator row
            if re.match(r'^\s*\|?\s*[:\-\s|]+\s*\|?\s*$', line):
                continue # Skip table delimiter row like |---|---|
            parts = [c.strip() for c in line.strip().split("|")]
            if parts and parts[0] == "":
                parts = parts[1:]
            if parts and parts[-1] == "":
                parts = parts[:-1]
            if parts:
                in_table = True
                table_rows.append(parts)
                continue
        else:
            if in_table:
                flush_table()

        if not stripped:
            continue

        # Headings
        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(stripped[2:])
            run.font.name = 'Calibri'
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[3:])
            run.font.name = 'Calibri'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        elif stripped.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(stripped[4:])
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)
        elif stripped.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(stripped[5:])
            run.font.name = 'Calibri'
            run.font.size = Pt(11.5)
            run.font.bold = True
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            clean_text = stripped[2:]
            # Clean basic bold markup
            parts = re.split(r'(\*\*.*?\*\*)', clean_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
        elif re.match(r'^\d+\.\s', stripped):
            match = re.match(r'^(\d+\.\s)(.*)', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(2)
            clean_text = match.group(2)
            parts = re.split(r'(\*\*.*?\*\*)', clean_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[2:])
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            parts = re.split(r'(\*\*.*?\*\*)', stripped)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)

    if in_table:
        flush_table()
    if in_code_block:
        flush_code_block()

    # Append embedded images if provided
    if embedded_images:
        doc.add_page_break()
        h = doc.add_heading("Visual Architecture & Technical Diagrams", level=1)
        h.paragraph_format.space_before = Pt(12)
        for img_path in embedded_images:
            if os.path.exists(img_path):
                doc.add_paragraph()
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(6)
                doc.add_picture(img_path, width=Inches(6.2))
                caption = doc.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption_run = caption.add_run(f"Figure: {os.path.basename(img_path)}")
                caption_run.font.size = Pt(9.5)
                caption_run.font.italic = True
                doc.add_paragraph()

    doc.save(docx_path)
    print(f"Created Word Document: {docx_path}")

# Run conversions
artifacts = [
    (
        "research/machine_learning_models_and_hardware_comparative_analysis.md",
        "research/machine_learning_models_and_hardware_comparative_analysis.docx",
        ["research/diagrams/ml_hardware_roofline.png", "research/diagrams/ai_ml_dl_nesting_dolls.png"]
    ),
    (
        "research/master_chief_ml_bootcamp_10th_grade.md",
        "research/master_chief_ml_bootcamp_10th_grade.docx",
        ["research/diagrams/spartan_ml_radar_clustering.png", "research/diagrams/ai_ml_dl_nesting_dolls.png"]
    ),
    (
        "research/ml_vs_deep_learning_connection.md",
        "research/ml_vs_deep_learning_connection.docx",
        ["research/diagrams/ai_ml_dl_nesting_dolls.png"]
    ),
    (
        "research/home_network_llm_mcp_comparative_analysis.md",
        "research/home_network_llm_mcp_comparative_analysis.docx",
        ["research/diagrams/master_chief_lan_topology.png", "research/diagrams/bland_voice_latency_pipeline.png"]
    ),
    (
        "research/master_chief_app_system_alignment.md",
        "research/master_chief_app_system_alignment.docx",
        ["research/diagrams/master_chief_lan_topology.png"]
    ),
    (
        "docs/chat_logs/session_ml_research_chat_log.md",
        "docs/chat_logs/session_ml_research_chat_log.docx",
        [
            "research/diagrams/ai_ml_dl_nesting_dolls.png",
            "research/diagrams/spartan_ml_radar_clustering.png",
            "research/diagrams/ml_hardware_roofline.png",
            "research/diagrams/master_chief_lan_topology.png",
            "research/diagrams/bland_voice_latency_pipeline.png"
        ]
    )
]

for md_file, docx_file, imgs in artifacts:
    markdown_to_docx(md_file, docx_file, imgs)

print("All documents converted to .docx successfully!")
